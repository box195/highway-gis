"""Read monthly source workbooks; write a private intermediate outside the repo.

Usage: python extract-inspection-history.py SOURCE_DIRECTORY PRIVATE_OUTPUT.json
Read only: source workbooks and their cached formula values remain unchanged.
"""
import calendar
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl


def text(value):
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def choose_sources(directory):
    chosen = {}
    for path in sorted(Path(directory).glob("*.xlsx")):
        match = re.match(r"(\d{1,2})월", path.name)
        if not match or path.name.startswith("~$"):
            continue
        month = int(match[1])
        rank = ("총괄" in path.stem, path.stat().st_mtime, path.name)
        if month not in chosen or rank > chosen[month][0]:
            chosen[month] = (rank, path)
    return [chosen[m][1] for m in sorted(chosen)]


def extract_sheet(sheet, year, month):
    branch = re.sub(r"[-_]통합.*$", "", sheet.title).strip()
    merged = {}
    for region in sheet.merged_cells.ranges:
        value = sheet.cell(region.min_row, region.min_col).value
        for row in range(region.min_row, region.max_row + 1):
            for col in range(region.min_col, min(region.max_col, 40) + 1):
                merged[row, col] = value

    def value(row, col):
        return sheet.cell(row, col).value or merged.get((row, col))

    groups = {}
    day_columns = {}
    section = ""
    target = ""
    road, milepost = {}, {}
    line = False
    invalid = []
    marked = 0
    for row in range(1, sheet.max_row + 1):
        actual = [sheet.cell(row, c).value for c in range(1, min(sheet.max_column, 45) + 1)]
        heading = text(actual[0])
        if "점검계획표" in heading:
            section = heading.replace("점검계획표", "").strip()
            day_columns = {}
            target = ""
            road, milepost = {}, {}
            continue
        numbers = {c + 1: v for c, v in enumerate(actual) if isinstance(v, int) and not isinstance(v, bool) and 1 <= v <= 31}
        if len(numbers) >= 28 and list(numbers.values())[:28] == list(range(1, 29)):
            day_columns = numbers
            target = ""
            line = "선로" in section
            continue
        if not day_columns or heading.startswith("*"):
            continue
        first_day = min(day_columns)
        left = [text(value(row, c)) for c in range(1, first_day)]
        if "노선" in left or "이정" in left:
            destination = road if "노선" in left else milepost
            for col in day_columns:
                destination[col] = text(value(row, col))
            continue
        if not line:
            current = text(value(row, 2))
            if current:
                target = current
            equipment = text(value(row, 5))
        else:
            equipment = " ".join(dict.fromkeys(s for s in (text(value(row, 2)), text(value(row, 4))) if s))
        for col, day in day_columns.items():
            raw = sheet.cell(row, col).value
            if raw is None or text(raw) == "":
                continue
            if day > calendar.monthrange(year, month)[1]:
                invalid.append({"sheet": sheet.title, "cell": sheet.cell(row, col).coordinate, "value": text(raw)})
                continue
            if not equipment or not (target or line):
                raise ValueError(f"Unclassified schedule cell: {sheet.title}!{sheet.cell(row,col).coordinate}")
            day_target = road.get(col, "선로") if line else target
            date = f"{year:04d}-{month:02d}-{day:02d}"
            span = milepost.get(col, "") if line else ""
            key = (date, section, day_target, span)
            group = groups.setdefault(key, {"date": date, "branch": branch, "section": section, "target": day_target, "milepost": span, "sheet": sheet.title, "items": []})
            group["items"].append({"equipment": equipment, "value": text(raw), "cycle": text(value(row, 5)) if line else "", "cell": sheet.cell(row,col).coordinate})
            marked += 1
    return list(groups.values()), marked, invalid


def extract(directory):
    months = []
    for path in choose_sources(directory):
        month = int(re.match(r"(\d+)월", path.name)[1])
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=False)
        years = set()
        for sheet in workbook:
            for row in sheet.iter_rows(max_row=6, max_col=5):
                for cell in row:
                    match = re.search(r"(20\d{2})년\s*(\d{1,2})월", text(cell.value))
                    if match:
                        if int(match[2]) != month:
                            raise ValueError(f"Month mismatch in {path.name}")
                        years.add(int(match[1]))
        if len(years) != 1:
            raise ValueError(f"Ambiguous workbook year: {path.name}")
        year = years.pop()
        groups, invalid, marked = [], [], 0
        for sheet in workbook:
            if re.sub(r"[-_]통합.*$", "", sheet.title).strip() not in {"동서울지사", "경기광주지사"}:
                continue
            rows, count, issues = extract_sheet(sheet, year, month)
            groups.extend(rows)
            marked += count
            invalid.extend(issues)
        workbook.close()
        if sum(len(g["items"]) for g in groups) != marked:
            raise ValueError("Source cell reconciliation failed")
        months.append({"month": f"{year:04d}-{month:02d}", "source": path.name, "source_kind": "총괄" if "총괄" in path.name else "1팀", "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "branches": sorted({g["branch"] for g in groups}), "groups": sorted(groups,key=lambda g:(g["date"],g["branch"],g["target"],g["section"])), "cell_count": marked, "invalid_dates": invalid})
    return {"version": 1, "months": months}


if __name__ == "__main__":
    output = Path(sys.argv[2]).resolve()
    repo = Path(__file__).resolve().parents[1]
    if output.is_relative_to(repo):
        raise SystemExit("Private workbook extraction must stay outside the repository")
    output.write_text(json.dumps(extract(sys.argv[1]),ensure_ascii=False),encoding="utf-8")
